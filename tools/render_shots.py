"""Render example outputs (xlsx sheets, docx, html) to PNG images for the workbook.

  uv run --with openpyxl --with pycel --with pillow python tools/render_shots.py OUTDIR

Spreadsheet images are drawn from the real files: cell values, fills, bold and column
widths come from openpyxl; formula cells are evaluated with pycel. A sheet-tab strip is
drawn along the bottom so it reads as a workbook.
"""
import html
import subprocess
import sys
import time
from pathlib import Path

import openpyxl
from PIL import Image, ImageChops

BASE = Path(__file__).resolve().parent.parent
EX = BASE / "example-outputs"
CHROME = "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome"
TMP = Path("/private/tmp/claude-501/claimfox-shots")
TMP.mkdir(parents=True, exist_ok=True)


def shoot(html_path, png_path, width=1400, height=2400):
    p = subprocess.Popen([CHROME, "--headless=new", "--disable-gpu", "--hide-scrollbars", "--force-device-scale-factor=2",
                          f"--screenshot={png_path}", f"--window-size={width},{height}", f"file://{html_path}"],
                         stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    for _ in range(40):
        time.sleep(0.5)
        if Path(png_path).exists() and Path(png_path).stat().st_size > 0:
            time.sleep(1.0)
            break
    p.kill()
    trim(png_path)


def trim(png_path, pad=24):
    im = Image.open(png_path).convert("RGB")
    bg = Image.new("RGB", im.size, im.getpixel((im.width - 2, im.height - 2)))
    box = ImageChops.difference(im, bg).getbbox()
    if box:
        l, t, r, b = box
        im = im.crop((0, 0, im.width, min(im.height, b + pad)))
    im.save(png_path, optimize=True)


def color(c):
    try:
        if c is not None and c.type == "rgb" and c.rgb and c.rgb not in ("00000000",):
            return "#" + c.rgb[-6:]
    except Exception:
        pass
    return None


def sheet_html(xlsx, sheet, max_rows=18, max_cols=8, title=None, col_width_scale=7.2):
    wb = openpyxl.load_workbook(xlsx)
    ws = wb[sheet]
    ev = None
    try:
        from pycel import ExcelCompiler
        ev = ExcelCompiler(filename=str(xlsx))
    except Exception:
        ev = None
    rows = []
    for r in range(1, min(ws.max_row, max_rows) + 1):
        cells, vals = [], []
        for c in range(1, min(ws.max_column, max_cols) + 1):
            cell = ws.cell(r, c)
            v = cell.value
            if isinstance(v, str) and v.startswith("=") and ev:
                try:
                    v = ev.evaluate(f"'{sheet}'!{cell.coordinate}")
                except Exception:
                    v = "#calc"
            if hasattr(v, "strftime"):
                v = v.strftime("%Y-%m-%d")
            if isinstance(v, int) and not isinstance(v, bool) and abs(v) >= 1000:
                v = f"{v:,}"
            if isinstance(v, float):
                fmt = cell.number_format or ""
                v = f"${v:,.2f}" if "$" in fmt else (f"{v:,.2f}" if "0.00" in fmt else (f"{v:,.0f}" if v == int(v) else f"{v:,.2f}"))
            txt = "" if v is None else html.escape(str(v))
            vals.append(txt.strip())
            bg = color(cell.fill.fgColor) if cell.fill and cell.fill.fill_type == "solid" else None
            fg = color(cell.font.color) if cell.font and cell.font.color else None
            style = []
            if bg:
                style.append(f"background:{bg}")
            if fg:
                style.append(f"color:{fg}")
            if cell.font and cell.font.bold:
                style.append("font-weight:600")
            if cell.font and cell.font.sz and cell.font.sz >= 13:
                style.append(f"font-size:{cell.font.sz + 2}px")
            cells.append(f'<td style="{";".join(style)}">{txt}</td>')
        filled = [i for i, x in enumerate(vals) if x]
        if filled == [0] and len(cells) > 1:
            cells = [cells[0].replace("<td ", f'<td colspan="{len(cells)}" ', 1)]
        rows.append("<tr>" + "".join(cells) + "</tr>")
    widths = []
    for c in range(1, min(ws.max_column, max_cols) + 1):
        letter = openpyxl.utils.get_column_letter(c)
        w = ws.column_dimensions[letter].width or 12
        widths.append(f'<col style="width:{min(w, 60) * col_width_scale:.0f}px">')
    tabs = "".join(f'<span class="tab{" on" if s == sheet else ""}">{html.escape(s)}</span>' for s in wb.sheetnames[:9])
    return f"""<!doctype html><meta charset="utf-8"><style>
body{{margin:0;background:#fff;font:13px/1.35 -apple-system,"Segoe UI",Helvetica,Arial,sans-serif;color:#1c2a39}}
.bar{{background:#217346;color:#fff;padding:8px 14px;font-size:13px}}
.wrap{{padding:0;overflow:hidden}}
table{{border-collapse:collapse;table-layout:fixed}} td[colspan]{{white-space:normal}}
td{{border:1px solid #e3e6ea;padding:5px 7px;vertical-align:top;white-space:normal;overflow:hidden;max-height:70px}}
.tabs{{border-top:1px solid #d0d4d9;background:#f3f4f6;padding:0 8px;display:flex;gap:2px}}
.tab{{padding:6px 12px;color:#57606a;font-size:12px}}.tab.on{{background:#fff;color:#217346;font-weight:600;border-bottom:2px solid #217346}}
</style><div class="bar">{html.escape(title or Path(xlsx).name)}</div><div class="wrap"><table>{"".join(widths)}{"".join(rows)}</table></div><div class="tabs">{tabs}</div>"""


def render_sheet(xlsx, sheet, out, width=1400, **kw):
    h = TMP / (Path(out).stem + ".html")
    h.write_text(sheet_html(xlsx, sheet, **kw))
    shoot(h, out, width=width)
    print("wrote", out)


def render_docx(docx, out, width=1000):
    h = TMP / (Path(out).stem + ".html")
    body = subprocess.run(["pandoc", str(docx), "-t", "html"], capture_output=True, text=True).stdout
    h.write_text(f"""<!doctype html><meta charset="utf-8"><style>
body{{margin:0;background:#e9ecef;font:14px/1.55 "Open Sans",Helvetica,Arial,sans-serif;color:#1c2a39}}
.page{{background:#fff;max-width:820px;margin:24px auto;padding:56px 64px;box-shadow:0 1px 4px rgba(0,0,0,.15)}}
h1{{color:#1c2a39;font-size:24px}} h2{{color:#1c2a39;font-size:18px;border-bottom:2px solid #F15A29;padding-bottom:4px}} h3{{font-size:15px}}
table{{border-collapse:collapse;width:100%;font-size:12.5px}} td,th{{border:1px solid #d9dde2;padding:5px 7px;vertical-align:top}} th{{background:#1c2a39;color:#fff}}
</style><div class="page">{body}</div>""")
    shoot(h, out, width=width, height=2200)
    print("wrote", out)


def render_html(src, out, width=1200, height=1800, crop_h=None):
    tmp_png = TMP / (Path(out).stem + "-full.png")
    shoot(Path(src), tmp_png, width=width, height=height)
    im = Image.open(tmp_png)
    if crop_h:
        im = im.crop((0, 0, im.width, min(im.height, crop_h * 2)))
    im.save(out, optimize=True)
    print("wrote", out)


if __name__ == "__main__":
    out = Path(sys.argv[1])
    out.mkdir(parents=True, exist_ok=True)
    B = EX / "Barbara - Contract requirements" / "Contract Requirements - Portfolio - 2026-09-28.xlsx"
    render_sheet(B, "Portfolio", out / "barbara-portfolio.png", max_rows=16, max_cols=4, width=1500)
    render_sheet(B, "Harborline Mutual", out / "barbara-matrix.png", max_rows=12, max_cols=7, width=1600, col_width_scale=6.2)
    C = EX / "Chris - DDQ responder" / "Pioneer Standard DDQ - Review Sheet - 2026-09-28.xlsx"
    render_sheet(C, "Review", out / "chris-review.png", max_rows=14, max_cols=7, width=1600, col_width_scale=6.0)
    render_sheet(C, "Questions for owners", out / "chris-owners.png", max_rows=12, max_cols=5, width=1400)
    G = EX / "Christine and Scott - GL review" / "GL Review - August 2026.xlsx"
    render_sheet(G, "Exceptions", out / "gl-exceptions.png", max_rows=12, max_cols=6, width=1600, col_width_scale=5.8)
    render_sheet(G, "Proof", out / "gl-proof.png", max_rows=14, max_cols=4, width=1500, col_width_scale=6.5)
    M = EX / "Mary - UAT builder" / "EREQ-2147 UAT - 2026-09-28.xlsx"
    render_sheet(M, "Summary", out / "mary-summary.png", max_rows=22, max_cols=5, width=1200)
    render_sheet(M, "Test cases", out / "mary-cases.png", max_rows=10, max_cols=9, width=1700, col_width_scale=5.5)
    render_sheet(M, "Due-date matrix", out / "mary-duedates.png", max_rows=12, max_cols=7, width=1400)
    K = EX / "Kaela - Onboarding builder" / "Onboarding Checklists - Global Intake Associate - 2026-10-05.xlsx"
    render_sheet(K, "Readiness gate", out / "kaela-gate.png", max_rows=8, max_cols=7, width=1400)
    render_sheet(K, "Missing info", out / "kaela-missing.png", max_rows=12, max_cols=6, width=1500)
    MI = EX / "Michelle - Requestor report cards" / "Requestor Report Cards - Q3 2026.xlsx"
    render_sheet(MI, "Archive watchlist", out / "michelle-watchlist.png", max_rows=12, max_cols=8, width=1500)
    render_docx(EX / "Kaela - Onboarding builder" / "Onboarding Plan - Global Intake Associate - 2026-10-05.docx", out / "kaela-plan.png")
    render_docx(EX / "Mary - UAT builder" / "EREQ-2147 UAT Test Plan - 2026-09-28.docx", out / "mary-plan.png")
    render_html(EX / "Amanda - Client review" / "Harborline Mutual Insurance - Quarterly Review - Q3 2026.html", out / "amanda-review.png", width=1100, height=1500)
    render_html(EX / "Michelle - Requestor report cards" / "Requestor Report Cards - Q3 2026.html", out / "michelle-card.png", width=1100, height=1150, crop_h=1100)
    render_html(EX / "Michelle - Requestor report cards" / "Statement of Open Invoices - RecordPoint Retrieval - 2026-09-29.html", out / "michelle-statement.png", width=1100, height=1200, crop_h=900)
    render_html(EX / "Fig - Weekly brief" / "Weekly Brief - week of 2026-09-21.html", out / "fig-brief.png", width=1100, height=1400, crop_h=1100)
