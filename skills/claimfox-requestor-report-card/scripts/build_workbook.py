#!/usr/bin/env python3
"""Turn a JSON spec into a clean, ClaimFox-branded Excel workbook.

  python3 build_workbook.py spec.json output.xlsx

Spec (every key except "sheets" is optional):
{
  "title": "Workbook title shown at the top of every sheet",
  "subtitle": "Source files, date prepared",
  "banner": "PROVISIONAL - Exhibit B not provided",        # orange warning line under the title
  "sheets": [
    {
      "name": "Sheet name (max 31 chars)",
      "columns": ["ID", "Item", "Amount", "Due", "Status", "Result"],
      "rows": [["A-1", "Example", 1250.5, "2026-10-09", "Needs review", ""]],
      "status_columns": ["Status"],                   # cells colored by value (see STATUS_COLORS)
      "status_colors": {"My value": "FFF2CC"},        # add or override colors
      "formats": {"Amount": "currency", "Due": "date", "Rate": "percent", "Count": "integer"},
      "dropdowns": {"Result": ["Pass", "Fail", "Blocked", "Not run"]},
      "widths": {"Item": 60},
      "row_height": 45                                 # optional fixed height for wrapped rows
    }
  ]
}
Header row: row 1 if no title; otherwise after the title lines plus one blank row (title+subtitle+banner = row 5).
Cells whose value starts with "=" are written as live Excel formulas, e.g. "=COUNTIF('Test cases'!K:K,\"Pass\")".
Dates as "YYYY-MM-DD" strings are converted to real Excel dates when the column has "date" format.
Requires openpyxl (`pip install openpyxl` if missing).
"""
import datetime as dt
import json
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

NAVY, ORANGE, SLATE = "1C2A39", "F15A29", "576374"
RED, AMBER, YELLOW, GREEN, BLUE, GREY = "F8D7D3", "FCE4D6", "FFF2CC", "E2EFDA", "DDEBF7", "E8EDF2"
STATUS_COLORS = {
    # severity / risk
    "high": RED, "s1": RED, "critical": RED, "medium": AMBER, "s2": AMBER, "low": GREY, "s3": YELLOW, "s4": GREY,
    # contract comparison
    "stricter": AMBER, "adds": BLUE, "looser": YELLOW, "matches": GREEN, "not specified": YELLOW, "unclear": RED,
    # answers
    "prior answer": GREEN, "drawn from policy": BLUE, "conflict": RED, "needs review": YELLOW,
    # test results
    "pass": GREEN, "fail": RED, "blocked": YELLOW, "not run": GREY,
    # generic
    "done": GREEN, "open": YELLOW, "late": RED, "due today": RED, "due this week": AMBER,
}
FORMATS = {"currency": '"$"#,##0.00;[Red]-"$"#,##0.00', "date": "yyyy-mm-dd", "percent": "0.0%",
           "integer": "#,##0", "number": "#,##0.00"}


def main(spec_path, out_path):
    spec = json.load(open(spec_path, encoding="utf-8"))
    wb = Workbook()
    wb.remove(wb.active)
    thin = Side(style="thin", color="D9DDE2")
    for sh in spec["sheets"]:
        ws = wb.create_sheet(sh["name"][:31])
        cols = sh["columns"]
        r = 1
        if spec.get("title"):
            ws.cell(r, 1, spec["title"]).font = Font(bold=True, size=14, color=NAVY)
            r += 1
        if spec.get("subtitle"):
            ws.cell(r, 1, spec["subtitle"]).font = Font(italic=True, size=10, color=SLATE)
            r += 1
        if spec.get("banner"):
            ws.cell(r, 1, spec["banner"]).font = Font(bold=True, size=11, color=ORANGE)
            r += 1
        r0 = r + 1 if r > 1 else 1
        for j, c in enumerate(cols, 1):
            cell = ws.cell(r0, j, c)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor=NAVY)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            cell.border = Border(bottom=Side(style="medium", color=ORANGE))
        colors = dict(STATUS_COLORS)
        colors.update({k.lower(): v for k, v in sh.get("status_colors", {}).items()})
        status_cols = sh.get("status_columns") or ([sh["status_column"]] if sh.get("status_column") else [])
        sidx = [cols.index(c) + 1 for c in status_cols if c in cols]
        fmts = {cols.index(c) + 1: FORMATS.get(f, f) for c, f in sh.get("formats", {}).items() if c in cols}
        for i, row in enumerate(sh["rows"], r0 + 1):
            for j, v in enumerate(row, 1):
                if j in fmts and fmts[j] == FORMATS["date"] and isinstance(v, str) and len(v) == 10:
                    try:
                        v = dt.date.fromisoformat(v)
                    except ValueError:
                        pass
                cell = ws.cell(i, j, v)
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                cell.border = Border(bottom=thin)
                if j in fmts:
                    cell.number_format = fmts[j]
                if i % 2 == 0:
                    cell.fill = PatternFill("solid", fgColor="FAFBFC")
            for j in sidx:
                if j - 1 < len(row):
                    val = str(row[j - 1]).strip().lower()
                    key = next((k for k in colors if val == k or val.startswith(k + " ")), None)
                    if key:
                        c = ws.cell(i, j)
                        c.fill = PatternFill("solid", fgColor=colors[key])
                        c.font = Font(bold=True, color=NAVY)
            if sh.get("row_height"):
                ws.row_dimensions[i].height = sh["row_height"]
        last = r0 + max(len(sh["rows"]), 1)
        for col, options in sh.get("dropdowns", {}).items():
            if col in cols:
                dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True)
                letter = get_column_letter(cols.index(col) + 1)
                dv.add(f"{letter}{r0 + 1}:{letter}{max(last, r0 + 200)}")
                ws.add_data_validation(dv)
        widths = sh.get("widths", {})
        for j, c in enumerate(cols, 1):
            longest = max([len(str(c))] + [len(str(rw[j - 1])) for rw in sh["rows"] if j - 1 < len(rw)])
            ws.column_dimensions[get_column_letter(j)].width = widths.get(c, min(max(10, longest + 2), 55))
        ws.freeze_panes = ws.cell(r0 + 1, 1)
        if sh["rows"]:
            ws.auto_filter.ref = f"A{r0}:{get_column_letter(len(cols))}{r0 + len(sh['rows'])}"
        ws.sheet_view.showGridLines = False
    wb.save(out_path)
    print(f"saved {out_path}")


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
