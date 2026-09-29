#!/usr/bin/env python3
"""ClaimFox month-end GL validation and bank reconciliation.

Deterministic checks. Claude runs this script and explains the results; Claude never does
the arithmetic itself. Standard library only, plus openpyxl for the Excel output if installed.

Usage:
  python3 validate_gl.py --gl gl_2026-08.csv --prior gl_2026-07.csv gl_2026-06.csv \
      --coa chart_of_accounts.csv --rules coding_rules.csv \
      [--bank bank_statement_2026-08.csv --disbursements ap_disbursements_2026-08.csv] \
      [--variance-pct 25 --variance-floor 1000 --round-threshold 10000 --near-dup-days 14 --split-limit 500] \
      --out ./gl-review-2026-08

Expected GL columns (rename yours to match first):
  entry_id (optional), date, journal_ref, account, account_name, vendor_or_customer, description,
  reference, debit, credit, source, entered_by, memo

Outputs in --out: exceptions.csv, account_variance.csv, summary.json and gl-review.xlsx with
Exceptions, Account variance, Checks run, Proof (live Excel formulas over the raw rows) and GL data tabs.
"""
import argparse
import csv
import datetime as dt
import json
import re
from collections import defaultdict
from pathlib import Path

REQUIRED = ["date", "journal_ref", "account", "vendor_or_customer", "description", "reference", "debit", "credit"]
VAGUE = re.compile(r"^\s*(adj(ustment)?|misc|per discussion|see me|tbd|reclass|correction|plug|true[- ]?up)\b", re.I)


def money(x):
    x = (x or "").replace(",", "").replace("$", "").strip()
    if x.startswith("(") and x.endswith(")"):
        x = "-" + x[1:-1]
    return round(float(x), 2) if x else 0.0


def fmt(x):
    return f"-${abs(x):,.2f}" if x < 0 else f"${x:,.2f}"


def load(path):
    with open(path, newline="", encoding="utf-8-sig") as f:
        rows = list(csv.DictReader(f))
    missing = [c for c in REQUIRED if rows and c not in rows[0]]
    if missing:
        raise SystemExit(f"{path}: missing columns {missing}. Rename the export's columns to the expected names first.")
    for i, r in enumerate(rows, start=2):
        r["_row"], r["_file"] = i, Path(path).name
        r["_dr"], r["_cr"] = money(r.get("debit")), money(r.get("credit"))
        r["_amt"] = r["_dr"] or r["_cr"]
        r["_date"] = dt.date.fromisoformat(r["date"][:10])
        r["_id"] = r.get("entry_id") or f'{r["_file"]}:{i}'
    return rows


CHECKS = {
    "Unbalanced entry": "Every journal entry's debits must equal its credits. Differences divisible by 9 usually mean swapped digits.",
    "Duplicate": "Same vendor, same invoice/reference, same amount and account posted in more than one journal entry.",
    "Possible duplicate": "Same vendor and amount (>= $500) with different references within {near_dup_days} days.",
    "Repeated identical charge": "Same vendor and exact amount 3+ times in the month (under $500). Can be legit (per-seat fees) or a billing loop.",
    "Split purchase pattern": "3+ charges from one vendor on the same day, or same-day charges that total more than ${split_limit} while each is under it.",
    "Coding": "Vendor matches a rule in coding_rules.csv but is posted to a different account.",
    "Missing recurring": "Vendor billed in every prior month provided but not this month.",
    "Variance": "Revenue or expense account moved at least ${variance_floor} and {variance_pct}% vs the prior month.",
    "Manual entry review": "Manual journal entry with 2+ of: weekend date, round amount >= ${round_threshold}, vague description, no reference.",
    "AP payments not in GL": "Vendor bills are credited to AP but no payments ever debit AP in the months provided.",
    "Check sequence gap": "Check numbers in the disbursements register skip a number.",
    "Payment does not match a bill": "A payment in the disbursements register has no vendor bill of the same amount in the months provided.",
    "Bank rec": "Every bank line matched to a GL cash entry or a register payment of the same amount within 7 days.",
}


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--gl", required=True)
    ap.add_argument("--prior", nargs="*", default=[])
    ap.add_argument("--coa")
    ap.add_argument("--rules")
    ap.add_argument("--bank")
    ap.add_argument("--disbursements")
    ap.add_argument("--variance-pct", type=float, default=25.0)
    ap.add_argument("--variance-floor", type=float, default=1000.0)
    ap.add_argument("--round-threshold", type=float, default=10000.0)
    ap.add_argument("--near-dup-days", type=int, default=14)
    ap.add_argument("--split-limit", type=float, default=500.0)
    ap.add_argument("--out", default="gl-review")
    a = ap.parse_args()
    S = {k: getattr(a, k) for k in ("variance_pct", "variance_floor", "round_threshold", "near_dup_days", "split_limit")}

    out = Path(a.out)
    out.mkdir(parents=True, exist_ok=True)
    gl = load(a.gl)
    priors = [load(p) for p in a.prior]
    coa = {}
    if a.coa:
        with open(a.coa, newline="", encoding="utf-8-sig") as f:
            coa = {r["account"]: r for r in csv.DictReader(f)}
    ex = []

    def flag(check, severity, rows, finding, amount=None, action=""):
        rows = rows if isinstance(rows, list) else [rows]
        ex.append({
            "check": check, "severity": severity,
            "journal_ref": ", ".join(sorted({r["journal_ref"] for r in rows if r.get("journal_ref")})),
            "entry_ids": ", ".join(r["_id"] for r in rows),
            "date": rows[0]["date"] if rows else "", "account": rows[0].get("account", "") if rows else "",
            "vendor": rows[0].get("vendor_or_customer", "") if rows else "",
            "amount": round(amount, 2) if amount is not None else "", "finding": finding, "suggested_action": action,
            "why_flagged": CHECKS[check.split(":")[0]].format(**S), "linked_findings": "",
        })

    by_je = defaultdict(list)
    for r in gl:
        by_je[r["journal_ref"]].append(r)

    # 1 balance
    unbalanced = []
    for je, rows in by_je.items():
        dsum, csum = round(sum(r["_dr"] for r in rows), 2), round(sum(r["_cr"] for r in rows), 2)
        if abs(dsum - csum) > 0.005:
            diff = round(abs(dsum - csum), 2)
            unbalanced.append(je)
            hint = " The difference is divisible by 9, which usually means two digits were swapped." if round(diff * 100) % 900 == 0 else ""
            flag("Unbalanced entry", "High", rows, f"{je}: debits {fmt(dsum)} vs credits {fmt(csum)}, off by {fmt(diff)}.{hint}", diff,
                 "Void and re-post the entry from the source document. Don't patch one side.")

    # 2 exact duplicates
    seen = defaultdict(list)
    for r in gl:
        if r["_dr"] and r.get("reference") and r.get("vendor_or_customer"):
            seen[(r["vendor_or_customer"].lower(), r["reference"].lower(), r["_dr"], r["account"])].append(r)
    dup_keys = set()
    for k, rows in seen.items():
        if len({r["journal_ref"] for r in rows}) > 1:
            dup_keys.add((k[0], k[2]))
            flag("Duplicate", "High", rows, f"{rows[0]['vendor_or_customer']} invoice {rows[0]['reference']} for {fmt(k[2])} is posted {len(rows)} times.",
                 k[2] * (len(rows) - 1), "Void the extra posting and confirm only one payment went out.")

    # 3 near duplicates
    groups = defaultdict(list)
    for r in gl:
        if r["_dr"] and r.get("vendor_or_customer") and not r["account"].startswith(("1", "2")):
            groups[(r["vendor_or_customer"].lower(), r["_dr"])].append(r)
    for (v, amt), rows in groups.items():
        if (v, amt) in dup_keys or len({r["journal_ref"] for r in rows}) < 2:
            continue
        rows.sort(key=lambda r: r["_date"])
        span = (rows[-1]["_date"] - rows[0]["_date"]).days
        if amt >= 500 and span <= a.near_dup_days and len({r.get("reference") for r in rows}) > 1:
            flag("Possible duplicate", "Medium", rows,
                 f"{rows[0]['vendor_or_customer']}: {len(rows)} charges of {fmt(amt)} within {span} days, different invoice numbers.",
                 amt, "Confirm these are separate services. Pull both invoices.")
        elif amt < 500 and len(rows) >= 3:
            flag("Repeated identical charge", "Low", rows,
                 f"{rows[0]['vendor_or_customer']}: {len(rows)} charges of exactly {fmt(amt)} this month ({fmt(amt * len(rows))} total).",
                 amt * len(rows), "Confirm the count matches seats or orders. Check it isn't also paid another way (e.g. an annual license).")

    # 3b split purchases
    day = defaultdict(list)
    for r in gl:
        if r["_dr"] and r.get("vendor_or_customer") and not r["account"].startswith(("1", "2")):
            day[(r["vendor_or_customer"], r["date"])].append(r)
    for (v, dte), rows in day.items():
        tot = round(sum(r["_dr"] for r in rows), 2)
        under = all(r["_dr"] < a.split_limit for r in rows)
        if len(rows) >= 3 or (len(rows) >= 2 and under and tot > a.split_limit):
            flag("Split purchase pattern", "Low", rows, f"{v}: {len(rows)} charges on {dte} totaling {fmt(tot)}"
                 + (f", each under the {fmt(a.split_limit)} limit." if under else "."), tot, "Check whether one purchase was split to stay under an approval limit.")

    # 4 coding
    if a.rules:
        with open(a.rules, newline="", encoding="utf-8-sig") as f:
            rules = [r for r in csv.DictReader(f) if r.get("vendor_pattern")]
        for r in gl:
            if not r["_dr"] or r["account"].startswith(("1", "2")):
                continue
            text = f'{r.get("vendor_or_customer", "")} {r.get("description", "")}'
            for rule in rules:
                if re.search(rule["vendor_pattern"], text, re.I):
                    allowed = [x.strip() for x in rule["expected_account"].split("|")]
                    if r["account"] not in allowed:
                        exp = " or ".join(f"{x} {coa.get(x, {}).get('account_name', '')}".strip() for x in allowed)
                        got = f"{r['account']} {coa.get(r['account'], {}).get('account_name', r.get('account_name', ''))}".strip()
                        flag("Coding", "Medium", r, f"{r['vendor_or_customer']} ({r['description']}) is coded to {got}. Rule says {exp}. {rule.get('note', '')}".strip(),
                             r["_dr"], f"Reclass to {allowed[0]} if the rule is right.")
                    break

    # 5 missing recurring
    if priors:
        cur_vendors = {r["vendor_or_customer"] for r in gl if r["_dr"]}
        prior_sets = [{r["vendor_or_customer"] for r in p if r["_dr"] and r.get("source") != "Card feed"} for p in priors]
        for v in sorted(set.intersection(*prior_sets) - cur_vendors):
            if not v:
                continue
            last = max((r for p in priors for r in p if r["vendor_or_customer"] == v and r["_dr"]), key=lambda r: r["_date"])
            flag("Missing recurring", "Medium", last,
                 f"{v} was billed in each of the prior {len(priors)} month(s) (last {fmt(last['_dr'])} on {last['date']}) but has nothing this month.",
                 last["_dr"], "Accrue the expected bill or confirm the service ended.")

    # 6 variance
    def totals(rows):
        t = defaultdict(float)
        for r in rows:
            if r["account"][:1] in ("4", "5", "6"):
                t[r["account"]] += r["_dr"] - r["_cr"]
        return t
    variances = []
    if priors:
        cur_t, pri_t = totals(gl), totals(priors[0])
        for acct in sorted(set(cur_t) | set(pri_t)):
            c, p = round(cur_t.get(acct, 0), 2), round(pri_t.get(acct, 0), 2)
            delta = round(c - p, 2)
            pctv = (delta / abs(p) * 100) if p else None
            name = coa.get(acct, {}).get("account_name", "")
            variances.append({"account": acct, "account_name": name, "prior": p, "current": c, "change": delta,
                              "change_pct": round(pctv, 1) if pctv is not None else ""})
            if abs(delta) >= a.variance_floor and (pctv is None or abs(pctv) >= a.variance_pct):
                drivers = sorted((r for r in gl if r["account"] == acct), key=lambda r: -r["_amt"])[:3]
                dtxt = "; ".join(f"{r['vendor_or_customer'] or r['description']} {fmt(r['_amt'])}" for r in drivers)
                txt = (f"{acct} {name} moved {fmt(delta)} ({pctv:+.1f}%) vs prior month: {fmt(p)} to {fmt(c)}. Biggest lines: {dtxt}."
                       if pctv is not None else f"{acct} {name} is new this month: {fmt(c)}.")
                flag("Variance", "Medium", drivers, txt, abs(delta), "Explain the driver or correct it.")
                ex[-1]["account"] = acct

    # 7 manual entries
    for je, rows in by_je.items():
        if not any((r.get("source") or "").lower().startswith("manual") for r in rows):
            continue
        amt = max(r["_amt"] for r in rows)
        reasons = []
        if rows[0]["_date"].weekday() >= 5:
            reasons.append(f"posted on a {rows[0]['_date'].strftime('%A')}")
        if amt >= a.round_threshold and amt % 1000 == 0:
            reasons.append(f"round amount {fmt(amt)}")
        memo = " ".join({(r.get("memo") or r.get("description") or "") for r in rows})
        if VAGUE.search(memo) or not memo.strip():
            reasons.append(f'vague description ("{memo.strip() or "blank"}")')
        if not any(r.get("reference") for r in rows):
            reasons.append("no supporting reference")
        if len(reasons) >= 2:
            flag("Manual entry review", "High" if len(reasons) >= 3 else "Medium", rows,
                 f"{je} by {rows[0].get('entered_by') or 'unknown'}: " + ", ".join(reasons) + ".", amt,
                 "Get the support and an approver other than the person who posted it.")

    # 8 AP relief
    months = [gl] + priors
    ap_cr = sum(r["_cr"] for m in months for r in m if r["account"] == "2000")
    ap_dr = sum(r["_dr"] for m in months for r in m if r["account"] == "2000")
    if ap_cr and not ap_dr:
        ex.append({"check": "AP payments not in GL", "severity": "Medium", "journal_ref": "", "entry_ids": "", "date": "", "account": "2000",
                   "vendor": "", "amount": round(ap_cr, 2), "finding": f"{fmt(ap_cr)} of vendor bills were credited to AP across {len(months)} month(s), but no payment ever debits AP. Either payments post somewhere else or AP is overstated.",
                   "suggested_action": "Find where vendor payments are recorded and tie AP to the payments register.", "why_flagged": CHECKS["AP payments not in GL"], "linked_findings": ""})

    # 9 bank rec + check sequence
    rec = None
    if a.disbursements:
        with open(a.disbursements, newline="", encoding="utf-8-sig") as f:
            disb = list(csv.DictReader(f))
        nums = sorted(int(m.group(1)) for d in disb for m in [re.search(r"CHK\s*(\d+)", d.get("check_or_ref", ""), re.I)] if m)
        for lo, hi in zip(nums, nums[1:]):
            for missing in range(lo + 1, hi):
                ex.append({"check": "Check sequence gap", "severity": "Medium", "journal_ref": "", "entry_ids": "", "date": "", "account": "1000",
                           "vendor": "", "amount": "", "finding": f"Check {missing} is missing from the register (between {lo} and {hi}).",
                           "suggested_action": "Confirm it was voided (and kept) or find where it went.", "why_flagged": CHECKS["Check sequence gap"], "linked_findings": ""})
        bills = [r for m in [gl] + priors for r in m if r["account"] == "2000" and r["_cr"]]
        for dd in disb:
            payee, amt = dd["payee"].lower(), money(dd["amount"])
            key = payee.split()[0]
            vb = [b for b in bills if key in (b.get("vendor_or_customer") or "").lower()]
            if vb and not any(abs(b["_cr"] - amt) < 0.005 for b in vb):
                amounts = sorted({b["_cr"] for b in vb})
                ex.append({"check": "Payment does not match a bill", "severity": "Medium", "journal_ref": "", "entry_ids": "", "date": dd["date"], "account": "2000",
                           "vendor": dd["payee"], "amount": amt, "finding": f'{dd["payee"]} {dd["check_or_ref"]} for {fmt(amt)} matches none of their bills in the period ({", ".join(fmt(x) for x in amounts[:4])}).',
                           "suggested_action": "Check whether it's a partial payment, a different invoice, or a keying error.", "why_flagged": CHECKS["Payment does not match a bill"], "linked_findings": ""})
    else:
        disb = []
    if a.bank:
        with open(a.bank, newline="", encoding="utf-8-sig") as f:
            bank = [dict(r, _amt=money(r["amount"]), _used=False) for r in csv.DictReader(f)]
        book = [{"date": r["date"], "desc": r["description"], "amt": round(r["_dr"] - r["_cr"], 2), "src": "GL cash (1000)"} for r in gl if r["account"] == "1000"]
        book += [{"date": d["date"], "desc": f'{d["payee"]} {d["check_or_ref"]}', "amt": -money(d["amount"]), "src": "disbursements register"} for d in disb]
        matched, unmatched = 0, []
        for b in book:
            hit = next((x for x in bank if not x["_used"] and abs(x["_amt"] - b["amt"]) < 0.005
                        and abs((dt.date.fromisoformat(x["date"]) - dt.date.fromisoformat(b["date"])).days) <= 7), None)
            if hit:
                hit["_used"] = True
                matched += 1
            else:
                unmatched.append(b)
        items = []
        for b in unmatched:
            num = re.search(r"CHK\s*\d+", b["desc"].upper())
            twin = next((x for x in bank if not x["_used"] and num and num.group(0).replace(" ", "") in x["description"].replace(" ", "")), None)
            if twin:
                twin["_used"] = True
                diff = round(abs(twin["_amt"]) - abs(b["amt"]), 2)
                note = " The difference is divisible by 9: likely swapped digits." if round(abs(diff) * 100) % 900 == 0 else ""
                items.append(("Amount mismatch", "High", abs(diff), f'{b["desc"]}: {b["src"]} shows {fmt(abs(b["amt"]))}, the bank cleared {fmt(abs(twin["_amt"]))} (difference {fmt(diff)}).{note}',
                              "Pull the cleared check image. Correct whichever record is wrong."))
            elif b["amt"] > 0:
                items.append(("Deposit in transit", "Low", b["amt"], f'{b["desc"]} ({b["src"]}, {b["date"]}) is not on the statement yet.', "Confirm it clears in the first days of next month."))
            else:
                items.append(("Outstanding payment", "Low", abs(b["amt"]), f'{b["desc"]} ({b["src"]}, {b["date"]}) has not cleared the bank.', "Normal timing. Follow up if still open after 30 days."))
        for x in bank:
            if not x["_used"]:
                items.append(("Not in books", "High", abs(x["_amt"]), f'{x["description"]} {fmt(x["_amt"])} on {x["date"]} is on the statement but not recorded anywhere.',
                              "Record it (bank fees go to 6800)." if "CHARGE" in x["description"] or "FEE" in x["description"] else "Find the source and record it."))
        for typ, sev, amt, txt, act in items:
            ex.append({"check": f"Bank rec: {typ}", "severity": sev, "journal_ref": "", "entry_ids": "", "date": "", "account": "1000", "vendor": "",
                       "amount": round(amt, 2), "finding": txt, "suggested_action": act, "why_flagged": CHECKS["Bank rec"], "linked_findings": ""})
        bank_net = round(sum(x["_amt"] for x in bank), 2)
        book_net = round(sum(b["amt"] for b in book), 2)
        adj = round(sum((-b["amt"] if b in unmatched else 0) for b in book), 2)
        rec = {"book_lines": len(book), "bank_lines": len(bank), "matched": matched, "bank_net_activity": bank_net, "book_net_activity": book_net,
               "note": "The statement has no opening or closing balance, so this is an activity reconciliation. Ask for the full statement to tie balances."}

    # link variances to other findings on the same account
    for e in ex:
        if e["check"] != "Variance":
            continue
        links = [f'{o["check"]} ({o["journal_ref"] or o["vendor"]}, {fmt(o["amount"]) if isinstance(o["amount"], float) else o["amount"]})'
                 for o in ex if o is not e and o["check"] not in ("Variance",) and o["account"] == e["account"]]
        if not links and e["account"]:
            links = [f'{o["check"]} ({o["journal_ref"]})' for o in ex if o is not e and o["journal_ref"] and
                     any(r["account"] == e["account"] for je in o["journal_ref"].split(", ") for r in by_je.get(je, []))]
        e["linked_findings"] = "; ".join(links) if links else "None. Still needs an explanation."

    order = {"High": 0, "Medium": 1, "Low": 2}
    ex.sort(key=lambda e: (order[e["severity"]], e["check"]))
    fields = ["severity", "check", "finding", "amount", "suggested_action", "linked_findings", "why_flagged", "journal_ref", "date", "account", "vendor", "entry_ids"]
    with open(out / "exceptions.csv", "w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=fields, extrasaction="ignore")
        w.writeheader()
        w.writerows(ex)
    if variances:
        with open(out / "account_variance.csv", "w", newline="") as f:
            w = csv.DictWriter(f, fieldnames=list(variances[0]))
            w.writeheader()
            w.writerows(variances)
    counts = {c: sum(1 for e in ex if e["check"].split(":")[0] == c) for c in CHECKS}
    summary = {
        "file": Path(a.gl).name, "lines_checked": len(gl), "journal_entries": len(by_je),
        "total_debits": round(sum(r["_dr"] for r in gl), 2), "total_credits": round(sum(r["_cr"] for r in gl), 2),
        "exceptions": len(ex), "by_severity": {s: sum(1 for e in ex if e["severity"] == s) for s in order},
        "checks_run": counts, "bank_rec": rec, "settings": S | {"prior_months": [Path(p).name for p in a.prior]},
    }
    summary["out_of_balance_by"] = round(summary["total_debits"] - summary["total_credits"], 2)

    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
        NAVY = "1C2A39"
        wb = Workbook()
        def header(ws, cols, widths=None):
            ws.append(cols)
            for c in ws[ws.max_row]:
                c.font, c.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor=NAVY)
            for i, wdt in enumerate(widths or [], 1):
                ws.column_dimensions[get_column_letter(i)].width = wdt
            ws.freeze_panes = ws.cell(ws.max_row + 1, 1)
        ws = wb.active
        ws.title = "Exceptions"
        header(ws, fields, (9, 22, 70, 12, 45, 45, 45, 14, 11, 8, 26, 26))
        fills = {"High": "F8D7D3", "Medium": "FCEBD2", "Low": "E8EDF2"}
        for e in ex:
            ws.append([e.get(k, "") for k in fields])
            ws.cell(ws.max_row, 1).fill = PatternFill("solid", fgColor=fills[e["severity"]])
            ws.cell(ws.max_row, 4).number_format = '"$"#,##0.00'
            for c in ws[ws.max_row]:
                c.alignment = Alignment(wrap_text=True, vertical="top")
        if variances:
            vs = wb.create_sheet("Account variance")
            header(vs, ["account", "account_name", "prior", "current", "change", "change_pct"], (9, 34, 14, 14, 14, 11))
            for v in variances:
                vs.append([v["account"], v["account_name"], v["prior"], v["current"], v["change"], v["change_pct"] if v["change_pct"] != "" else None])
                for c in vs[vs.max_row][2:5]:
                    c.number_format = '#,##0.00;[Red]-#,##0.00'
        cr = wb.create_sheet("Checks run")
        header(cr, ["Check", "What it looks for", "Findings"], (26, 90, 10))
        for c, desc in CHECKS.items():
            cr.append([c, desc.format(**S), counts[c]])
            for cell in cr[cr.max_row]:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        # raw data tabs so the Proof tab can use live formulas
        cols = ["entry_id", "journal_ref", "date", "account", "account_name", "vendor_or_customer", "description", "reference", "debit", "credit", "source", "entered_by", "memo"]
        def raw(name, rows):
            s = wb.create_sheet(name)
            header(s, cols, (18, 11, 11, 8, 26, 26, 34, 18, 12, 12, 12, 12, 20))
            for r in rows:
                s.append([r["_id"], r["journal_ref"], r["date"], int(r["account"]) if r["account"].isdigit() else r["account"], r.get("account_name", ""), r.get("vendor_or_customer", ""),
                          r.get("description", ""), r.get("reference", ""), r["_dr"] or None, r["_cr"] or None, r.get("source", ""), r.get("entered_by", ""), r.get("memo", "")])
            return s
        pf = wb.create_sheet("Proof", 1)
        raw("GL data", gl)
        if priors:
            raw("Prior month GL data", priors[0])
        G, P = "'GL data'", "'Prior month GL data'"
        header(pf, ["What", "Live formula result", "Script result", "Formula"], (60, 18, 18, 70))
        def prow(label, formula, script):
            pf.append([label, formula, script, formula])
            pf.cell(pf.max_row, 4).value = "'" + formula  # show the formula text
            for c in (2, 3):
                pf.cell(pf.max_row, c).number_format = '#,##0.00;[Red]-#,##0.00'
        pf.append(["Every figure below is an Excel formula over the raw rows in the GL data tabs. Click any cell in column B to see it."])
        pf.cell(pf.max_row, 1).font = Font(italic=True, color="576374")
        prow("Total debits, this month", f"=SUM({G}!I:I)", summary["total_debits"])
        prow("Total credits, this month", f"=SUM({G}!J:J)", summary["total_credits"])
        prow("Out of balance by", f"=SUM({G}!I:I)-SUM({G}!J:J)", summary["out_of_balance_by"])
        for je in unbalanced:
            dsum = round(sum(r["_dr"] for r in by_je[je]), 2) - round(sum(r["_cr"] for r in by_je[je]), 2)
            prow(f"{je}: debits minus credits", f'=SUMIF({G}!B:B,"{je}",{G}!I:I)-SUMIF({G}!B:B,"{je}",{G}!J:J)', round(dsum, 2))
        for e in ex:
            if e["check"] == "Duplicate":
                rows = [r for r in gl if r["_id"] in e["entry_ids"].split(", ")]
                r = rows[0]
                prow(f"Times {r['vendor_or_customer']} ref {r['reference']} is posted as a debit",
                     f'=COUNTIFS({G}!F:F,"{r["vendor_or_customer"]}",{G}!H:H,"{r["reference"]}",{G}!I:I,{r["_dr"]})', len(rows))
        if priors:
            for e in ex:
                if e["check"] == "Variance":
                    acct = e["account"]
                    v = next((x for x in variances if x["account"] == acct), None)
                    if v is None:
                        continue
                    prow(f"{acct} change vs prior month", f'=(SUMIF({G}!D:D,"{acct}",{G}!I:I)-SUMIF({G}!D:D,"{acct}",{G}!J:J))-(SUMIF({P}!D:D,"{acct}",{P}!I:I)-SUMIF({P}!D:D,"{acct}",{P}!J:J))', v["change"])
        wb.save(out / "gl-review.xlsx")
        summary["excel"] = str(out / "gl-review.xlsx")
    except ImportError:
        summary["excel"] = None
    (out / "summary.json").write_text(json.dumps(summary, indent=2))
    print(json.dumps(summary, indent=2))


if __name__ == "__main__":
    main()
